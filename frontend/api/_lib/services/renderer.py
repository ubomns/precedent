import re
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ─────────────────────────────────────────────────────────────────────────────
# ID helpers (match frontend getDisplayId / normalizeId)
# ─────────────────────────────────────────────────────────────────────────────
_ID_PREFIXES = {
    "functional_requirements": "FR",
    "non_functional_requirements": "NFR",
    "acceptance_criteria": "AC",
    "user_access_management": "UAM",
    "security_controls": "SC",
    "benefits_realization": "BR",
}


def _normalize_id(raw: str) -> str:
    if not raw:
        return ""
    m = re.match(r"^([A-Za-z]+)-(\d+)$", raw)
    if not m:
        return raw
    num = int(m.group(2))
    return f"{m.group(1)}-{str(num).zfill(2)}"


def _display_id(section_key: str, item: dict, index: int) -> str:
    prefix = _ID_PREFIXES.get(section_key)
    if prefix:
        return f"{prefix}-{str(index + 1).zfill(2)}"
    raw = item.get("id") or item.get("requirement_id")
    if raw:
        return _normalize_id(raw)
    return ""


# ─────────────────────────────────────────────────────────────────────────────
# Style helpers
# ─────────────────────────────────────────────────────────────────────────────
_ID_COL_WIDTH = Inches(0.55)


def _set_run(run, font_name="Calibri", font_size=Pt(10), bold=False, color=RGBColor(0x33, 0x33, 0x33)):
    run.font.name = font_name
    run.font.size = font_size
    run.font.bold = bold
    run.font.color.rgb = color


def _add_page_break(doc):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    fldChar = OxmlElement("w:br")
    fldChar.set(qn("w:type"), "page")
    run._r.append(fldChar)


def _style_table_header(row):
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for run in paragraph.runs:
                _set_run(run, bold=True, color=RGBColor(0x1A, 0x1A, 0x1A))


def _set_cell_text(cell, text: str):
    """Replace cell text and apply justify + body font styling."""
    cell.text = text
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for run in paragraph.runs:
            _set_run(run)


def _set_narrow_id(cell, text: str):
    """Set a narrow ID column."""
    cell.width = _ID_COL_WIDTH
    _set_cell_text(cell, text)


# ─────────────────────────────────────────────────────────────────────────────
# Main render
# ─────────────────────────────────────────────────────────────────────────────
def render_brd_to_docx(brd: dict, buffer, project_name: str = ""):
    doc = Document()

    # ── Cover page ──
    # Main title (centered, large)
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run("Business Requirements Document")
    _set_run(run, font_size=Pt(24), bold=True, color=RGBColor(0x1A, 0x1A, 0x1A))

    doc.add_paragraph()
    doc.add_paragraph()

    # Title field
    title_label = doc.add_paragraph()
    run_label = title_label.add_run("Title: ")
    _set_run(run_label, bold=True, font_size=Pt(12), color=RGBColor(0x1A, 0x1A, 0x1A))
    run_val = title_label.add_run(project_name)
    _set_run(run_val, font_size=Pt(12), color=RGBColor(0x1A, 0x1A, 0x1A))

    # Date field
    date_para = doc.add_paragraph()
    run_label = date_para.add_run("Date: ")
    _set_run(run_label, bold=True, font_size=Pt(12), color=RGBColor(0x1A, 0x1A, 0x1A))
    run_val = date_para.add_run(datetime.now().strftime("%B %d, %Y"))
    _set_run(run_val, font_size=Pt(12), color=RGBColor(0x1A, 0x1A, 0x1A))

    # Page break before sections
    _add_page_break(doc)

    sections_order = [
        ("executive_summary", "01", "Executive Summary"),
        ("business_objective", "02", "Business Objective"),
        ("scope", "03", "Scope"),
        ("stakeholders", "04", "Stakeholders"),
        ("functional_requirements", "05", "Functional Requirements"),
        ("non_functional_requirements", "06", "Non-Functional Requirements"),
        ("assumptions", "07", "Assumptions"),
        ("constraints", "08", "Constraints"),
        ("acceptance_criteria", "09", "Acceptance Criteria"),
        ("user_access_management", "10", "User Access Management"),
        ("security_controls", "11", "Security Controls"),
        ("benefits_realization", "12", "Benefits Realization"),
    ]

    for key, num, heading in sections_order:
        section = brd.get(key)
        if not section:
            continue

        content = section.get("content", {})

        # Section Heading: "01  Executive Summary"
        h = doc.add_heading(f"{num}  {heading}", level=1)
        _set_run(h.runs[0], font_size=Pt(18), bold=True, color=RGBColor(0x1A, 0x1A, 0x1A))

        if key in ("executive_summary", "business_objective"):
            narrative = content.get("narrative", "")
            if isinstance(narrative, list):
                for para in narrative:
                    p = doc.add_paragraph(para)
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            else:
                p = doc.add_paragraph(narrative)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        elif key == "scope":
            p = doc.add_paragraph()
            run = p.add_run("In Scope")
            _set_run(run, bold=True, font_size=Pt(11))
            for item in content.get("in_scope", []):
                p = doc.add_paragraph(item, style="List Bullet")
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            p = doc.add_paragraph()
            run = p.add_run("Out of Scope")
            _set_run(run, bold=True, font_size=Pt(11))
            for item in content.get("out_of_scope", []):
                p = doc.add_paragraph(item, style="List Bullet")
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        elif key == "stakeholders":
            table = doc.add_table(rows=1, cols=3)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            _set_cell_text(hdr[0], "Name")
            _set_cell_text(hdr[1], "Role")
            _set_cell_text(hdr[2], "RACI")
            _style_table_header(table.rows[0])
            for s in content.get("stakeholders", []):
                row = table.add_row().cells
                _set_cell_text(row[0], s.get("name", ""))
                _set_cell_text(row[1], s.get("role", ""))
                _set_cell_text(row[2], s.get("raci", ""))

        elif key == "functional_requirements":
            items = content.get("requirements", [])
            table = doc.add_table(rows=1, cols=3)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            _set_narrow_id(hdr[0], "ID")
            _set_cell_text(hdr[1], "Requirement")
            _set_cell_text(hdr[2], "Priority")
            _style_table_header(table.rows[0])
            for i, r in enumerate(items):
                row = table.add_row().cells
                _set_narrow_id(row[0], _display_id(key, r, i))
                _set_cell_text(row[1], r.get("requirement", ""))
                _set_cell_text(row[2], r.get("priority", ""))

        elif key == "non_functional_requirements":
            items = content.get("requirements", [])
            table = doc.add_table(rows=1, cols=4)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            _set_narrow_id(hdr[0], "ID")
            _set_cell_text(hdr[1], "Category")
            _set_cell_text(hdr[2], "Requirement")
            _set_cell_text(hdr[3], "Metric")
            _style_table_header(table.rows[0])
            for i, r in enumerate(items):
                row = table.add_row().cells
                _set_narrow_id(row[0], _display_id(key, r, i))
                _set_cell_text(row[1], r.get("category", ""))
                _set_cell_text(row[2], r.get("requirement", ""))
                _set_cell_text(row[3], r.get("metric", ""))

        elif key == "assumptions":
            for item in content.get("assumptions", []):
                p = doc.add_paragraph(item, style="List Bullet")
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        elif key == "constraints":
            for item in content.get("constraints", []):
                p = doc.add_paragraph(item, style="List Bullet")
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        elif key == "acceptance_criteria":
            items = content.get("criteria", [])
            table = doc.add_table(rows=1, cols=2)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            _set_narrow_id(hdr[0], "ID")
            _set_cell_text(hdr[1], "Criteria")
            _style_table_header(table.rows[0])
            for i, c in enumerate(items):
                row = table.add_row().cells
                _set_narrow_id(row[0], _display_id(key, c, i))
                _set_cell_text(row[1], c.get("criteria", ""))

        elif key == "user_access_management":
            items = content.get("roles", [])
            table = doc.add_table(rows=1, cols=4)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            _set_narrow_id(hdr[0], "ID")
            _set_cell_text(hdr[1], "Role")
            _set_cell_text(hdr[2], "Permissions")
            _set_cell_text(hdr[3], "IAM Group")
            _style_table_header(table.rows[0])
            for i, r in enumerate(items):
                row = table.add_row().cells
                _set_narrow_id(row[0], _display_id(key, r, i))
                _set_cell_text(row[1], r.get("role", ""))
                _set_cell_text(row[2], r.get("permissions", ""))
                _set_cell_text(row[3], r.get("iam_group", ""))

        elif key == "security_controls":
            items = content.get("controls", [])
            table = doc.add_table(rows=1, cols=4)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            _set_narrow_id(hdr[0], "ID")
            _set_cell_text(hdr[1], "Domain")
            _set_cell_text(hdr[2], "Control")
            _set_cell_text(hdr[3], "Implementation")
            _style_table_header(table.rows[0])
            for i, c in enumerate(items):
                row = table.add_row().cells
                _set_narrow_id(row[0], _display_id(key, c, i))
                _set_cell_text(row[1], c.get("domain", ""))
                _set_cell_text(row[2], c.get("control", ""))
                _set_cell_text(row[3], c.get("implementation", ""))

        elif key == "benefits_realization":
            items = content.get("benefits", [])
            table = doc.add_table(rows=1, cols=4)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            _set_narrow_id(hdr[0], "ID")
            _set_cell_text(hdr[1], "Benefit")
            _set_cell_text(hdr[2], "Measurement")
            _set_cell_text(hdr[3], "Timeline")
            _style_table_header(table.rows[0])
            for i, b in enumerate(items):
                row = table.add_row().cells
                _set_narrow_id(row[0], _display_id(key, b, i))
                _set_cell_text(row[1], b.get("benefit", ""))
                _set_cell_text(row[2], b.get("measurement", ""))
                _set_cell_text(row[3], b.get("timeline", ""))

        doc.add_paragraph()

    doc.save(buffer)
